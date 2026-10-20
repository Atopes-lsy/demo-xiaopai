# -*- coding: utf-8 -*-
"""Generate the readable DOCX version of the 12-question survey."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "大学生课表与日程管理调查（12题）.docx",
)

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(4)


def add_para(
    text,
    bold=False,
    size=11,
    align=None,
    space_before=0,
    space_after=4,
    indent=None,
    color=None,
):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = RGBColor(*color)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p


add_para(
    "大学生课表与日程管理习惯调查",
    bold=True,
    size=18,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_after=2,
)
add_para(
    "用户研究问卷 · 12 题 · 约 2 分钟",
    size=10,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_after=10,
    color=(0x66, 0x66, 0x66),
)

intro = (
    "同学你好！本问卷想了解大学生如何管理课表、日程和群里的通知，约需 2 分钟。"
    "匿名填写，数据仅用于研究。提交后可参与红包抽奖（共 10 个 1 元微信红包），"
    "审核通过后发放。感谢你的帮助！"
)
add_para(intro, size=10.5, space_after=10)

questions = [
    (
        "1. 您目前是否在校大学生？",
        ["A. 是，本科生", "B. 是，研究生", "C. 否，非在校学生"],
    ),
    (
        "2. 您的年级是？",
        ["A. 大一", "B. 大二", "C. 大三", "D. 大四", "E. 研究生及以上"],
    ),
    (
        "3. 您目前在用哪些工具管理课表和日程？（多选）",
        [
            "A. 学校官方课表小程序或教务系统",
            "B. 超级课程表、WakeUp 等课表 App",
            "C. 手机自带日历或备忘录",
            "D. 滴答清单、Todoist 等待办 App",
            "E. 截图当壁纸或纸质课表",
            "F. 微信收藏或文件传输助手",
            "G. 其他",
        ],
    ),
    (
        "4. 您会在课表或日历工具里，添加课程以外的安排吗（班会、讲座、志愿活动、面试等）？",
        ["A. 经常", "B. 偶尔", "C. 很少", "D. 从不"],
    ),
    (
        "5. 过去一个月，您大约有多少次需要记录课程以外的安排（班会、活动、作业截止等）？",
        ["A. 0 次", "B. 1-2 次", "C. 3-5 次", "D. 6-10 次", "E. 10 次以上"],
    ),
    (
        "6. 收到一条需要记录的通知（例如“周四下午 3 点在 A301 开年级会”），您通常怎么处理？",
        [
            "A. 立刻记到日历或备忘录",
            "B. 截图保存",
            "C. 记到课表工具里",
            "D. 转发到微信收藏或文件传输助手",
            "E. 只看一眼，靠脑子记",
            "F. 不处理",
            "G. 其他",
        ],
    ),
    (
        "7. 过去一个月，您是否因为忘记或漏看通知，错过或差点错过某件事？",
        [
            "A. 是，发生过多次",
            "B. 是，发生过 1-2 次",
            "C. 没有，但差点发生",
            "D. 完全没有",
        ],
    ),
    (
        "8. 您最容易漏掉或记错的信息是？（多选）",
        [
            "A. 时间",
            "B. 地点",
            "C. 截止日期",
            "D. 需要携带或准备的东西",
            "E. 活动名称或具体内容",
            "F. 很少漏掉",
        ],
    ),
    (
        "9. 您通常如何区分“要去参加的活动”和“必须在某时前完成的 DDL”？",
        [
            "A. 分开管理，放在不同地方",
            "B. 放在一起，但自己用标记区分",
            "C. 不区分，都记在一起",
            "D. 我没有同时管理这两类事",
        ],
    ),
    (
        "10. 您希望课表和课程外日程在同一屏查看吗？",
        [
            "A. 非常希望",
            "B. 比较希望",
            "C. 一般",
            "D. 不太需要",
            "E. 完全不需要",
        ],
    ),
    (
        "11. 假设一个微信小助手只能看到你主动转发的那条消息，不会读取聊天记录。"
        "你愿意把群通知（文字或截图）转发给它，让它自动识别成日程或 DDL 吗？",
        [
            "A. 非常愿意",
            "B. 比较愿意",
            "C. 一般",
            "D. 不太愿意",
            "E. 完全不愿意",
        ],
    ),
    (
        "12. 识别后需要您确认或修改，您最多能接受多长时间？",
        [
            "A. 5 秒以内",
            "B. 10 秒以内",
            "C. 20 秒以内",
            "D. 只要结果准确，多花一点时间也可以",
            "E. 不愿意确认，最好直接存好",
        ],
    ),
]

for stem, options in questions:
    add_para(stem, bold=True, size=11.5, space_before=8, space_after=3)
    for option in options:
        add_para(option, size=11, indent=0.75, space_after=1)

add_para(
    "问卷结束，感谢你的参与！",
    size=10.5,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    space_before=14,
)

footer = doc.sections[0].footer.paragraphs[0]
footer.text = "大学生课表与日程管理习惯调查"
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in footer.runs:
    run.font.size = Pt(9)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

doc.save(OUT)
print("saved:", OUT)
